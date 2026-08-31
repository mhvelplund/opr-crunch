"""The generated formulas must agree with the reference model.

This is the check that the spreadsheet translation in opr_crunch.formulas did
not drift from opr_crunch.model. It evaluates the real formula strings written
onto the Tests tab, using the same subset of functions a spreadsheet would.
"""

import pytest
from openpyxl.utils import get_column_letter

from opr_crunch.build import INPUT_HEADERS, build_workbook
from opr_crunch.formulas import CALC_STEPS
from opr_crunch.model import resolve
from opr_crunch.scenarios import TEST_SCENARIOS

from .formula_eval import SheetEvaluator

STEP_NAMES = [name for name, _ in CALC_STEPS]
FIRST_STEP_COL = 2 + len(INPUT_HEADERS)


@pytest.fixture(scope="module")
def tests_sheet():
    return build_workbook()["Tests"]


def step_address(step: str, row: int) -> str:
    return f"{get_column_letter(FIRST_STEP_COL + STEP_NAMES.index(step))}{row}"


@pytest.mark.parametrize(
    "name,scenario",
    TEST_SCENARIOS,
    ids=[name for name, _ in TEST_SCENARIOS],
)
def test_formula_hits_match_the_model(tests_sheet, name, scenario):
    row = 2 + [n for n, _ in TEST_SCENARIOS].index(name)
    got = SheetEvaluator(tests_sheet).value(step_address("hits", row))
    assert got == pytest.approx(float(resolve(scenario).hits))


@pytest.mark.parametrize(
    "name,scenario",
    TEST_SCENARIOS,
    ids=[name for name, _ in TEST_SCENARIOS],
)
def test_formula_wounds_match_the_model(tests_sheet, name, scenario):
    row = 2 + [n for n, _ in TEST_SCENARIOS].index(name)
    got = SheetEvaluator(tests_sheet).value(step_address("wounds", row))
    assert got == pytest.approx(float(resolve(scenario).wounds))


def test_every_scenario_row_is_marked_pass(tests_sheet):
    """The Result column is what the user reads after importing the sheet."""
    evaluator = SheetEvaluator(tests_sheet)
    result_col = get_column_letter(
        1 + len(INPUT_HEADERS) + len(STEP_NAMES) + 5
    )
    for i in range(len(TEST_SCENARIOS)):
        row = 2 + i
        assert evaluator.value(f"{result_col}{row}") == "PASS", TEST_SCENARIOS[i][0]
