"""Structure of the generated workbook."""

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from opr_crunch import csvdata, layout
from opr_crunch.build import build_workbook
from opr_crunch.layout import SIMULATION_COLUMNS, SIM_FIRST_ROW, SIM_ROWS, last_row
from opr_crunch.model import MODELLED_RULES, valid_model_counts

from .formula_eval import WorkbookEvaluator


@pytest.fixture(scope="module")
def wb():
    return build_workbook()


def headers(sheet):
    return [c.value for c in sheet[1]]


class TestTabs:
    def test_simulation_is_the_leftmost_tab(self, wb):
        assert wb.sheetnames[0] == "Simulation"

    def test_all_tabs_are_present(self, wb):
        assert wb.sheetnames == [
            "Simulation",
            "Units",
            "Weapons",
            "Upgrades",
            "Special rules",
            "Factions",
            "Lists",
            "Calc",
            "Tests",
        ]

    @pytest.mark.parametrize("name", ["Lists", "Calc", "Tests"])
    def test_working_tabs_are_hidden(self, wb, name):
        assert wb[name].sheet_state == "hidden"

    @pytest.mark.parametrize(
        "name", ["Simulation", "Units", "Weapons", "Upgrades", "Special rules", "Factions"]
    )
    def test_content_tabs_are_visible(self, wb, name):
        assert wb[name].sheet_state == "visible"


class TestSimulationTab:
    def test_headers_match_the_spec(self, wb):
        assert headers(wb["Simulation"]) == SIMULATION_COLUMNS

    def test_special_rules_columns_are_last(self, wb):
        assert SIMULATION_COLUMNS[-2:] == ["Att. Special Rules", "Def. Special Rules"]

    def test_hits_and_wounds_precede_special_rules(self, wb):
        assert SIMULATION_COLUMNS[-5:-2] == ["Hits", "Wounds", "Cost/Wound"]

    def test_cost_per_wound_follows_wounds(self, wb):
        assert SIMULATION_COLUMNS[SIMULATION_COLUMNS.index("Wounds") + 1] == "Cost/Wound"

    def test_every_row_has_a_hits_formula(self, wb):
        sim = wb["Simulation"]
        col = SIMULATION_COLUMNS.index("Hits") + 1
        for row in range(SIM_FIRST_ROW, SIM_FIRST_ROW + SIM_ROWS):
            assert str(sim.cell(row=row, column=col).value).startswith("=")

    def test_hits_delegates_to_the_calc_tab(self, wb):
        sim = wb["Simulation"]
        col = SIMULATION_COLUMNS.index("Hits") + 1
        assert "Calc!" in sim.cell(row=SIM_FIRST_ROW, column=col).value


class TestDataTabs:
    def test_units_headers(self, wb):
        assert headers(wb["Units"]) == [
            "Name", "Faction", "Quality", "Defense", "Size", "Tough", "Ravage", "Impact",
            "Special Rules", "Cost", "Key",
        ]

    def test_weapons_headers(self, wb):
        assert headers(wb["Weapons"]) == [
            "Name", "Faction", "Unit", "Range", "Attacks", "AP", "Blast", "Deadly", "Impact",
            "Special Rules", "Cost", "Key",
        ]

    def test_upgrades_headers(self, wb):
        assert headers(wb["Upgrades"]) == [
            "Name", "Faction", "Unit", "AP", "Blast", "Deadly", "Impact", "Special Rules",
            "Cost", "Key",
        ]

    def test_factions_are_seeded(self, wb):
        names = [r[0].value for r in wb["Factions"].iter_rows(min_row=2) if r[0].value]
        assert "Robot Legions" in names

    def test_every_modelled_rule_is_documented(self, wb):
        """The tab is the user's reference for what the sheet actually computes."""
        rows = [r for r in wb["Special rules"].iter_rows(min_row=2) if r[0].value]
        listed = {r[0].value for r in rows}
        assert set(MODELLED_RULES) <= listed

    def test_modelled_rules_are_flagged_as_such(self, wb):
        cols = headers(wb["Special rules"])
        name_i, flag_i = cols.index("Name"), cols.index("Modelled")
        for row in wb["Special rules"].iter_rows(min_row=2):
            if not row[name_i].value:
                continue
            expected = "Yes" if row[name_i].value in MODELLED_RULES else "No"
            assert row[flag_i].value == expected, row[name_i].value

    def test_reference_only_rules_are_listed_too(self, wb):
        """Scraped army-book rules the engine ignores still belong in the tab."""
        listed = {r[0].value for r in wb["Special rules"].iter_rows(min_row=2) if r[0].value}
        assert "Infiltrate" in listed and "Infiltrate" not in MODELLED_RULES

    @pytest.mark.parametrize("tab", ["Units", "Weapons", "Upgrades"])
    def test_lookup_keys_are_unique(self, wb, tab):
        """Names repeat across army books and across units; the full key must not.

        Several books print a CCW or a Stomp, so lookups match on Name and
        Faction. Weapons and Upgrades add Unit too, since the same weapon
        name can belong to more than one unit as a separate row.
        """
        cols = headers(wb[tab])
        indices = [cols.index(part) for part in layout.KEY_PARTS[tab]]
        keys = [
            tuple(r[i].value for i in indices)
            for r in wb[tab].iter_rows(min_row=2)
            if r[0].value
        ]
        assert len(keys) == len(set(keys)), tab

    def test_a_name_shared_by_two_factions_is_kept_distinct(self, wb):
        rows = [
            (r[0].value, r[1].value)
            for r in wb["Weapons"].iter_rows(min_row=2)
            if r[0].value
        ]
        shared = {n for n, f in rows if sum(1 for m, _ in rows if m == n) > 1}
        assert shared, "expected at least one name shared across factions"


class TestDropdownCapacity:
    """The cascading dropdowns scan a fixed range; the data must fit inside it.

    Importing an army book grows these tabs, and a row past the end of the
    range simply never appears in the dropdown - silently.
    """

    @pytest.mark.parametrize(
        "tab,loader",
        [
            ("Units", csvdata.units),
            ("Weapons", csvdata.weapons),
            ("Upgrades", csvdata.upgrades),
            ("Factions", csvdata.factions),
        ],
    )
    def test_every_row_is_reachable_from_a_dropdown(self, tab, loader):
        last_used = len(loader()) + 1  # +1 for the header row
        assert last_used <= last_row(tab), (
            f"{tab} reaches row {last_used} but dropdowns only scan to "
            f"{last_row(tab)}; raise layout.CAPACITY[{tab!r}]"
        )


class TestSeededRowsResolve:
    """Importing a faction deletes its old rows, which can orphan a seeded pick.

    A Simulation row naming a unit or weapon that no longer exists still looks
    fine - the lookups just return nothing and the row quietly reads zero.
    """

    @pytest.mark.parametrize(
        "column,loader",
        [
            ("Attacker", csvdata.units),
            ("Defender", csvdata.units),
            ("Weapon", csvdata.weapons),
            ("Att. upgrades", csvdata.upgrades),
            ("Def. upgrades", csvdata.upgrades),
        ],
    )
    def test_every_seeded_pick_exists(self, column, loader):
        known = {row[0] for row in loader()}
        index = SIMULATION_COLUMNS.index(column)
        for number, row in enumerate(csvdata.simulation(), start=2):
            picked = row[index]
            assert not picked or picked in known, (
                f"csv/simulation.csv row {number}: {column}={picked!r} "
                f"is not in the data"
            )


class TestSeededModelCounts:
    """A seeded count must be legal for the unit it is paired with.

    Re-importing a book can change a unit's size, which would leave the
    shipped workbook opening with rows already flagged red.
    """

    @pytest.mark.parametrize(
        "unit_column,count_column",
        [("Attacker", "#"), ("Defender", "Def. #")],
    )
    def test_every_seeded_count_is_legal(self, unit_column, count_column):
        sizes = {(row[0], row[1]): row[4] for row in csvdata.units()}
        unit_i = SIMULATION_COLUMNS.index(unit_column)
        count_i = SIMULATION_COLUMNS.index(count_column)
        faction_i = SIMULATION_COLUMNS.index(
            "Att. faction" if unit_column == "Attacker" else "Def. faction"
        )

        for number, row in enumerate(csvdata.simulation(), start=2):
            if not row[unit_i]:
                continue
            size = sizes[(row[unit_i], row[faction_i])]
            allowed = valid_model_counts(size)
            assert row[count_i] in allowed, (
                f"csv/simulation.csv row {number}: {count_column}={row[count_i]} "
                f"but {row[unit_i]} is size {size}, so only {allowed} are legal"
            )


class TestValidation:
    def test_dropdowns_are_attached_to_the_simulation_tab(self, wb):
        assert len(wb["Simulation"].data_validations.dataValidation) > 0

    def test_faction_dropdown_reads_from_the_factions_tab(self, wb):
        sources = [dv.formula1 for dv in wb["Simulation"].data_validations.dataValidation]
        assert any("Factions" in str(s) for s in sources)

    def test_unit_dropdowns_read_from_the_lists_tab(self, wb):
        sources = [dv.formula1 for dv in wb["Simulation"].data_validations.dataValidation]
        assert any("Lists" in str(s) for s in sources)


class TestSheetProtection:
    """Simulation is the only tab meant for direct editing; every other tab
    is either sourced from csv/ (edit + rebuild instead) or internal
    plumbing, so it's protected against accidental edits.
    """

    def test_simulation_is_not_protected(self, wb):
        assert wb["Simulation"].protection.sheet is False

    @pytest.mark.parametrize(
        "name",
        ["Units", "Weapons", "Upgrades", "Special rules", "Factions", "Lists", "Calc", "Tests"],
    )
    def test_other_tabs_are_protected(self, wb, name):
        assert wb[name].protection.sheet is True


class TestStaleDependentPicks:
    """Changing a faction or unit can't clear the now-invalid pick beneath it
    (no Apps Script), so those cells are flagged with conditional formatting
    instead - one rule per cell, keyed off the same filtered Lists range the
    cell's own dropdown already uses.
    """

    @pytest.mark.parametrize(
        "header", ["Attacker", "Att. upgrades", "Weapon", "Defender", "Def. upgrades"]
    )
    def test_every_dependent_cell_has_a_stale_check(self, wb, header):
        sim = wb["Simulation"]
        col = get_column_letter(SIMULATION_COLUMNS.index(header) + 1)
        ranges = {str(rng.sqref) for rng in sim.conditional_formatting._cf_rules}
        for row in range(SIM_FIRST_ROW, SIM_FIRST_ROW + SIM_ROWS):
            assert f"{col}{row}" in ranges

    def test_stale_check_compares_against_the_row_s_own_filtered_list(self, wb):
        sim = wb["Simulation"]
        col = get_column_letter(SIMULATION_COLUMNS.index("Attacker") + 1)
        rules_by_range = {
            str(rng.sqref): rules for rng, rules in sim.conditional_formatting._cf_rules.items()
        }
        formula = rules_by_range[f"{col}{SIM_FIRST_ROW}"][0].formula[0]
        units_last_row = last_row("Units")
        assert formula == (
            f'AND(${col}{SIM_FIRST_ROW}<>"",'
            f'COUNTIF(Lists!$A$2:$A${units_last_row},${col}{SIM_FIRST_ROW})=0)'
        )


class TestModelCountDropdowns:
    """A unit is fielded at its listed size or double it, so the count is a
    dropdown of those values rather than a free number - and, like every other
    dependent pick, it is flagged when the unit changes underneath it.
    """

    @pytest.mark.parametrize("header", ["#", "Def. #"])
    def test_the_count_is_a_dropdown(self, wb, header):
        sim = wb["Simulation"]
        col = get_column_letter(SIMULATION_COLUMNS.index(header) + 1)
        targets = {
            ref
            for dv in sim.data_validations.dataValidation
            for ref in str(dv.sqref).split()
        }
        for row in range(SIM_FIRST_ROW, SIM_FIRST_ROW + SIM_ROWS):
            assert f"{col}{row}" in targets

    @pytest.mark.parametrize("header", ["#", "Def. #"])
    def test_an_illegal_count_is_flagged(self, wb, header):
        sim = wb["Simulation"]
        col = get_column_letter(SIMULATION_COLUMNS.index(header) + 1)
        ranges = {str(rng.sqref) for rng in sim.conditional_formatting._cf_rules}
        for row in range(SIM_FIRST_ROW, SIM_FIRST_ROW + SIM_ROWS):
            assert f"{col}{row}" in ranges

    @pytest.mark.parametrize(
        "unit,offered",
        [
            ("Warriors", [5, 10]),      # size 5, so 5 or the doubled 10
            ("Robot Snakes", [3, 6]),   # doubling follows the unit's own size
            ("Robot Lord", [1, ""]),    # one model, and no doubled form
        ],
    )
    def test_the_dropdown_offers_the_size_and_double_it(self, unit, offered):
        """A single-model unit's second option blanks itself out, and a blank
        cell is omitted from the dropdown rather than offered as an option.
        """
        wb = build_workbook()
        sim = wb["Simulation"]
        for header, value in (("Att. faction", "Robot Legions"), ("Attacker", unit)):
            sim.cell(row=SIM_FIRST_ROW, column=SIMULATION_COLUMNS.index(header) + 1, value=value)

        source = next(
            str(dv.formula1)
            for dv in sim.data_validations.dataValidation
            if f"D{SIM_FIRST_ROW}" in str(dv.sqref).split()
        )
        assert source.startswith("Lists!")
        column = source.split("!$")[1].split("$")[0]
        evaluate = WorkbookEvaluator(wb, "Lists")
        assert [evaluate.cell("Lists", f"{column}{r}") for r in (2, 3)] == offered


class TestCostPerWound:
    """Points spent per projected wound, from the attacker's own cost."""

    def test_every_row_has_a_formula(self, wb):
        sim = wb["Simulation"]
        col = SIMULATION_COLUMNS.index("Cost/Wound") + 1
        for row in range(SIM_FIRST_ROW, SIM_FIRST_ROW + SIM_ROWS):
            assert str(sim.cell(row=row, column=col).value).startswith("=")

    def test_it_divides_by_the_averaged_wounds(self, wb):
        """Not a per-branch step: cost/avg(wounds) is not avg(cost/wounds)."""
        sim = wb["Simulation"]
        col = SIMULATION_COLUMNS.index("Cost/Wound") + 1
        wounds_col = SIMULATION_COLUMNS.index("Wounds") + 1
        formula = sim.cell(row=SIM_FIRST_ROW, column=col).value
        wounds = sim.cell(row=SIM_FIRST_ROW, column=wounds_col).value
        shared = wounds[wounds.index("Calc!"):].rstrip(")")
        assert shared in formula

    def test_no_projected_wounds_shows_infinity(self, wb):
        sim = wb["Simulation"]
        col = SIMULATION_COLUMNS.index("Cost/Wound") + 1
        assert "∞" in sim.cell(row=SIM_FIRST_ROW, column=col).value

    def test_an_empty_row_stays_blank(self, wb):
        sim = wb["Simulation"]
        col = SIMULATION_COLUMNS.index("Cost/Wound") + 1
        attacker = get_column_letter(SIMULATION_COLUMNS.index("Attacker") + 1)
        formula = sim.cell(row=SIM_FIRST_ROW, column=col).value
        assert f'${attacker}{SIM_FIRST_ROW}=""' in formula


class TestTestsTab:
    def test_has_a_pass_fail_column(self, wb):
        assert "Result" in headers(wb["Tests"])

    def test_carries_scenarios(self, wb):
        rows = [r for r in wb["Tests"].iter_rows(min_row=2) if r[0].value]
        assert len(rows) >= 10

    def test_compares_expected_against_a_live_formula(self, wb):
        tests = wb["Tests"]
        cols = headers(tests)
        live = tests.cell(row=2, column=cols.index("Actual hits") + 1).value
        assert str(live).startswith("=")
