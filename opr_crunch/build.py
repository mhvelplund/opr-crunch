"""Assemble dist/opr_crunch.xlsx."""

from __future__ import annotations

from itertools import zip_longest
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from opr_crunch import csvdata, layout
from opr_crunch.formulas import CALC_STEPS, Refs, has
from opr_crunch.model import Scenario, resolve, variable_effect
from opr_crunch.scenarios import TEST_SCENARIOS

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OUTPUT_FILL = PatternFill("solid", fgColor="E2EFDA")

#: Inputs the calc block reads. On Calc these are lookups into the data tabs;
#: on Tests they are literals. Either way the step formulas are identical.
INPUT_HEADERS = [
    "models", "attacks", "quality", "defense", "def_models",
    "ap", "blast", "deadly", "impact", "ravage", "tough", "att_rules", "def_rules",
    "var_active", "w_ap",
    "is_melee", "is_charging", "is_fatigued", "long_range", "in_cover",
]

#: The subset the step formulas read. var_active and w_ap drive the weighted
#: sum of the two branches instead, so they sit outside the block.
REF_INPUTS = [h for h in INPUT_HEADERS if h not in ("var_active", "w_ap")]

LAST_SIM_ROW = layout.SIM_FIRST_ROW + layout.SIM_ROWS - 1
SRC_FIRST = layout.SOURCE_FIRST_ROW

#: Generated Name|Faction column on each data tab. Lookups match on it rather
#: than on the name, because army books share generic names across factions.
KEY_COL = {
    "Units": get_column_letter(layout.key_column(layout.UNITS_COLUMNS)),
    "Weapons": get_column_letter(layout.key_column(layout.WEAPONS_COLUMNS)),
    "Upgrades": get_column_letter(layout.key_column(layout.UPGRADES_COLUMNS)),
}


# --------------------------------------------------------------------------
# small helpers
# --------------------------------------------------------------------------

def write_headers(ws: Worksheet, headers: List[str], width: int = 16) -> None:
    for i, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def _key_expr(name_ref: str, faction_ref: str, unit_ref: str = "") -> str:
    parts = [name_ref, faction_ref] + ([unit_ref] if unit_ref else [])
    return '&"|"&'.join(parts)


def lookup(
    sheet: str, value_col: str, name_ref: str, faction_ref: str,
    default: str = "0", unit_ref: str = "",
) -> str:
    """INDEX/MATCH on the tab's key.

    Matching on the name alone would be ambiguous: a CCW or a Stomp exists in
    several army books, and the wrong faction's profile would be returned.

    Weapons and Upgrades key on Name|Faction|Unit instead, since the same
    weapon name can belong to more than one unit with a different profile;
    those are written as separate rows. ``unit_ref`` is matched first, then
    falls back to a blank-Unit row (gear generic to the whole faction).
    """
    key = KEY_COL[sheet]
    index = f"INDEX('{sheet}'!${value_col}:${value_col},MATCH({{key}},'{sheet}'!${key}:${key},0))"
    specific = index.format(key=_key_expr(name_ref, faction_ref, unit_ref))
    if not unit_ref:
        return f"IFERROR({specific},{default})"
    generic = index.format(key=_key_expr(name_ref, faction_ref, '""'))
    return f"IFERROR({specific},IFERROR({generic},{default}))"


#: The merged column that stands in for both the old "Melee state" and
#: "Range >9\"" inputs - only one of the two ever applies to a given weapon,
#: since Combat is now derived from it.
STATE_COL = 'Melee state / Range >9"'


def textjoin(delimiter: str, ignore_empty: bool, *parts: str) -> str:
    """TEXTJOIN, prefixed for the OOXML "future function" convention.

    openpyxl writes formula text verbatim; functions added after the base
    Excel 2007 spec (TEXTJOIN among them) must be stored as ``_xlfn.NAME`` in
    the file or readers - Excel and Google Sheets alike - reject the name and
    show #NAME? instead of evaluating it.
    """
    flag = "TRUE" if ignore_empty else "FALSE"
    return f'_xlfn.TEXTJOIN("{delimiter}",{flag},' + ",".join(parts) + ")"


def sim(col: str, row: int) -> str:
    """Absolute-column, relative-row reference into the Simulation tab."""
    return f"Simulation!${get_column_letter(layout.col_index(col))}{row}"


def list_validation(source: str) -> DataValidation:
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    dv.showDropDown = False  # openpyxl inverts this flag; False means "show it"
    return dv


def inline_validation(options: List[str]) -> DataValidation:
    return list_validation('"' + ",".join(options) + '"')


# --------------------------------------------------------------------------
# data tabs
# --------------------------------------------------------------------------

def _faction_dropdown(ws: Worksheet, column: int, tab: str) -> None:
    dv = list_validation(f"Factions!$A${SRC_FIRST}:$A${layout.last_row('Factions')}")
    ws.add_data_validation(dv)
    letter = get_column_letter(column)
    dv.add(f"{letter}{SRC_FIRST}:{letter}{layout.last_row(tab)}")


def build_data_tab(
    wb: Workbook,
    title: str,
    columns: List[str],
    rows: List[list],
    width: int = 16,
    widths: Dict[str, int] | None = None,
    faction_column: int | None = None,
    wrap_column: int | None = None,
) -> None:
    """Write one CSV-sourced tab: header, rows, widths and faction dropdown."""
    ws = wb.create_sheet(title)
    write_headers(ws, columns, width=width)
    for letter, size in (widths or {}).items():
        ws.column_dimensions[letter].width = size

    for i, row in enumerate(rows, start=2):
        for j, value in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=value)
            if j == wrap_column:
                cell.alignment = Alignment(wrap_text=True)

    if faction_column:
        _faction_dropdown(ws, faction_column, title)

    if title in KEY_COL:
        key = layout.key_column(columns)
        cell = ws.cell(row=1, column=key, value="Key")
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        letters = [get_column_letter(columns.index(part) + 1) for part in layout.KEY_PARTS[title]]
        for i in range(2, len(rows) + 2):
            formula = '&"|"&'.join(f"${letter}{i}" for letter in letters)
            ws.cell(row=i, column=key, value=f"={formula}")
        ws.column_dimensions[get_column_letter(key)].width = 28


#: (title, columns, loader, kwargs) for every tab sourced from csv/.
DATA_TABS = [
    ("Units", layout.UNITS_COLUMNS, csvdata.units,
     dict(widths={"A": 22, "H": 34}, faction_column=2)),
    ("Weapons", layout.WEAPONS_COLUMNS, csvdata.weapons,
     dict(width=11, widths={"A": 22, "J": 34}, faction_column=2)),
    ("Upgrades", layout.UPGRADES_COLUMNS, csvdata.upgrades,
     dict(width=11, widths={"A": 22, "H": 34}, faction_column=2)),
    ("Special rules", layout.SPECIAL_RULES_COLUMNS, csvdata.special_rules,
     dict(width=14, widths={"A": 16, "D": 90}, wrap_column=4)),
    ("Factions", layout.FACTIONS_COLUMNS, csvdata.factions, dict(width=22)),
]


# --------------------------------------------------------------------------
# cascading dropdown source lists
# --------------------------------------------------------------------------

#: (label, source tab, Simulation column holding the faction to filter by,
#: Simulation column holding the unit to filter by, or "" for tabs with no
#: Unit column). A blank Unit cell in the source tab means "every unit", so
#: it always matches too - that's how faction-generic gear stays selectable.
LIST_SPECS = [
    ("Attacker units", "Units", "Att. faction", ""),
    ("Attacker upgrades", "Upgrades", "Att. faction", "Attacker"),
    ("Weapons", "Weapons", "Att. faction", "Attacker"),
    ("Defender units", "Units", "Def. faction", ""),
    ("Defender upgrades", "Upgrades", "Def. faction", "Defender"),
]


def build_lists(wb: Workbook) -> Dict[int, Dict[str, str]]:
    """One filtered name column per Simulation row per dropdown.

    Deliberately built from plain IF formulas rather than FILTER: the result is
    a fixed range that data validation can point at, and it survives the .xlsx
    round-trip into Google Sheets unchanged. Non-matching cells are left blank,
    which Sheets omits from the dropdown.
    """
    ws = wb.create_sheet("Lists")
    ranges: Dict[int, Dict[str, str]] = {}

    for i, row in enumerate(range(layout.SIM_FIRST_ROW, LAST_SIM_ROW + 1)):
        ranges[row] = {}
        for k, (label, source, faction_col, unit_col) in enumerate(LIST_SPECS):
            column = i * len(LIST_SPECS) + k + 1
            letter = get_column_letter(column)
            ws.cell(row=1, column=column, value=f"{label} r{row}")
            faction_key = sim(faction_col, row)
            last = layout.last_row(source)
            if unit_col:
                unit_key = sim(unit_col, row)
                source_columns = layout.UPGRADES_COLUMNS if source == "Upgrades" else layout.WEAPONS_COLUMNS
                unit_col_letter = get_column_letter(source_columns.index("Unit") + 1)
                condition = (
                    f"AND('{source}'!$B{{src}}={faction_key},"
                    f"OR('{source}'!${unit_col_letter}{{src}}={unit_key},"
                    f"'{source}'!${unit_col_letter}{{src}}=\"\"))"
                )
            else:
                condition = f"'{source}'!$B{{src}}={faction_key}"
            for src in range(SRC_FIRST, last + 1):
                cond = condition.format(src=src)
                ws.cell(
                    row=src,
                    column=column,
                    value=f"=IF({cond},'{source}'!$A{src},\"\")",
                )
            ranges[row][label] = f"Lists!${letter}${SRC_FIRST}:${letter}${last}"
    return ranges


#: Per-slot (melee state, range) pairing for the merged column's option list.
#: Combat="Melee" shows the melee-state name; Combat="Ranged" shows Yes/No
#: for the two slots YES_NO covers and nothing for the rest. Either way only
#: the options relevant to the row's derived Combat ever appear - Sheets
#: omits blank cells from a list-validation dropdown.
STATE_SLOTS = list(zip_longest(layout.MELEE_STATE_OPTIONS, layout.YES_NO, fillvalue=""))


def build_state_lists(wb: Workbook, start_column: int) -> Dict[int, str]:
    """One filtered options column per Simulation row for STATE_COL."""
    ws = wb["Lists"]
    ranges: Dict[int, str] = {}
    last = SRC_FIRST + len(STATE_SLOTS) - 1

    for i, row in enumerate(range(layout.SIM_FIRST_ROW, LAST_SIM_ROW + 1)):
        column = start_column + i
        letter = get_column_letter(column)
        ws.cell(row=1, column=column, value=f"Melee state/Range r{row}")
        combat = sim("Combat", row)
        for offset, (melee_val, range_val) in enumerate(STATE_SLOTS, start=SRC_FIRST):
            range_expr = f'"{range_val}"' if range_val else '""'
            ws.cell(
                row=offset,
                column=column,
                value=(
                    f'=IF({combat}="Melee","{melee_val}",'
                    f'IF({combat}="Ranged",{range_expr},""))'
                ),
            )
        ranges[row] = f"Lists!${letter}${SRC_FIRST}:${letter}${last}"
    return ranges


# --------------------------------------------------------------------------
# calc block
# --------------------------------------------------------------------------

def calc_inputs_for(row: int) -> List[str]:
    """Formulas pulling one Simulation row's inputs out of the data tabs."""
    weapon, att_up = sim("Weapon", row), sim("Att. upgrades", row)
    defender, def_up = sim("Defender", row), sim("Def. upgrades", row)
    combat, state = sim("Combat", row), sim(STATE_COL, row)
    attacker = sim("Attacker", row)
    att_fac, def_fac = sim("Att. faction", row), sim("Def. faction", row)
    blank = chr(34) * 2

    def pair(weapon_col: str, upgrade_col: str) -> str:
        return (
            f"=MAX({lookup('Weapons', weapon_col, weapon, att_fac, unit_ref=attacker)},"
            f"{lookup('Upgrades', upgrade_col, att_up, att_fac, unit_ref=attacker)})"
        )

    def input_cell(name: str) -> str:
        return f"${get_column_letter(INPUT_HEADERS.index(name) + 1)}{row}"

    att = input_cell("att_rules")
    melee = f'({combat}="Melee")'
    far = f'({state}="Yes")'
    fighter = has(att, "Unpredictable Fighter")
    versatile = f'AND({has(att, "Versatile Attack")},{far})'

    return [
        f"={sim('#', row)}",
        f"={lookup('Weapons', 'E', weapon, att_fac, unit_ref=attacker)}",
        f"={lookup('Units', 'C', attacker, att_fac)}",
        f"={lookup('Units', 'D', defender, def_fac)}",
        f"={sim('Def. #', row)}",
        f"={lookup('Weapons', 'F', weapon, att_fac, unit_ref=attacker)}"
        f"+{lookup('Upgrades', 'D', att_up, att_fac, unit_ref=attacker)}",
        pair("G", "E"),
        pair("H", "F"),
        pair("I", "G"),
        f"={lookup('Units', 'G', attacker, att_fac)}",
        f"={lookup('Units', 'F', defender, def_fac)}",
        # The attacking unit's own profile grants rules too (Relentless,
        # Piercing Assault, ...), so it joins the weapon's and the upgrade's.
        # TEXTJOIN drops blank parts, so this also reads cleanly as the
        # Simulation tab's "Att./Def. Special Rules" columns, not just an
        # internal has()-searched blob.
        "=" + textjoin(
            ", ", True,
            lookup("Units", "H", attacker, att_fac, blank),
            lookup("Weapons", "J", weapon, att_fac, blank, unit_ref=attacker),
            lookup("Upgrades", "H", att_up, att_fac, blank, unit_ref=attacker),
        ),
        "=" + textjoin(
            ", ", True,
            lookup("Units", "H", defender, def_fac, blank),
            lookup("Upgrades", "H", def_up, def_fac, blank, unit_ref=defender),
        ),
        # Rules granting "either AP(+1) or +1 to hit". Unpredictable Fighter
        # contains Unpredictable, so the longer name is excluded explicitly.
        f"=IF(OR(AND({fighter},{melee}),"
        f"AND({has(att, 'Unpredictable')},NOT({fighter})),"
        f"{versatile}),1,0)",
        f"=IF({input_cell('var_active')}=0,1,"
        f'IF({versatile},IF({sim("Variable effect", row)}="To-hit",0,1),0.5))',
        f"={melee}",
        f'=AND({combat}="Melee",{state}="Charging")',
        f'=AND({combat}="Melee",{state}="Fatigued")',
        f"={far}",
        f'=({sim("In cover", row)}="Yes")',
    ]


STEP_NAMES = [name for name, _ in CALC_STEPS]


def input_letters(row: int, first_col: int) -> Dict[str, str]:
    return {
        name: f"${get_column_letter(first_col + offset)}{row}"
        for offset, name in enumerate(INPUT_HEADERS)
    }


def write_calc_block(
    ws: Worksheet, row: int, first_col: int, step_col: int, extra_ap: str, extra_hit: str
) -> Dict[str, str]:
    """Emit one branch of the resolution and return each step's cell address."""
    letters = input_letters(row, first_col)
    refs = Refs(
        extra_ap=extra_ap,
        extra_hit=extra_hit,
        **{name: letters[name] for name in REF_INPUTS},
    )

    cells: Dict[str, str] = {}
    for offset, (name, builder) in enumerate(CALC_STEPS):
        column = step_col + offset
        address = f"${get_column_letter(column)}{row}"
        ws.cell(row=row, column=column, value=builder(refs, cells))
        cells[name] = address
    return cells


def write_both_branches(ws: Worksheet, row: int, first_col: int) -> Dict[str, str]:
    """Emit both branches and average them by the AP branch's weight.

    Rules like Unpredictable grant "either AP(+1) or +1 to hit", which cannot
    be folded into one pass. When no such rule is in play var_active is 0, both
    branches collapse onto the plain resolution and the weight is 1.
    """
    letters = input_letters(row, first_col)
    active, weight = letters["var_active"], letters["w_ap"]

    step_col = first_col + len(INPUT_HEADERS)
    ap_branch = write_calc_block(ws, row, first_col, step_col, active, "0")
    hit_col = step_col + len(CALC_STEPS)
    hit_branch = write_calc_block(ws, row, first_col, hit_col, "0", active)

    final = hit_col + len(CALC_STEPS)
    out = {}
    for offset, name in enumerate(("hits", "wounds")):
        column = final + offset
        ws.cell(
            row=row,
            column=column,
            value=f"={weight}*{ap_branch[name]}+(1-{weight})*{hit_branch[name]}",
        )
        out[name] = f"${get_column_letter(column)}{row}"
    return out


def calc_headers() -> List[str]:
    return (
        INPUT_HEADERS
        + [f"ap:{n}" for n in STEP_NAMES]
        + [f"hit:{n}" for n in STEP_NAMES]
        + ["hits", "wounds"]
    )


#: The weapon's own profile fields (not merged with any upgrade), each
#: computed once per row on the Calc tab so the Simulation tab's derived
#: Combat and Weapon stats columns can reference a plain cell instead of
#: repeating the same INDEX/MATCH lookup for every condition and value.
WEAPON_STAT_FIELDS = [
    ("range", "D", "0"),
    ("attacks", "E", "0"),
    ("ap", "F", "0"),
    ("blast", "G", "0"),
    ("deadly", "H", "0"),
    ("impact", "I", "0"),
    ("rules", "J", '""'),
]


def build_weapon_stat_calc(wb: Workbook, start_col: int) -> Dict[int, Dict[str, str]]:
    ws = wb["Calc"]
    for offset, (name, _, _) in enumerate(WEAPON_STAT_FIELDS):
        cell = ws.cell(row=1, column=start_col + offset, value=f"weapon:{name}")
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT

    cols: Dict[int, Dict[str, str]] = {}
    for row in range(layout.SIM_FIRST_ROW, LAST_SIM_ROW + 1):
        weapon = sim("Weapon", row)
        att_fac = sim("Att. faction", row)
        attacker = sim("Attacker", row)
        cols[row] = {}
        for offset, (name, letter, default) in enumerate(WEAPON_STAT_FIELDS):
            column = start_col + offset
            formula = f"={lookup('Weapons', letter, weapon, att_fac, default, unit_ref=attacker)}"
            ws.cell(row=row, column=column, value=formula)
            cols[row][name] = f"Calc!${get_column_letter(column)}{row}"
    return cols


def build_calc(wb: Workbook) -> Dict[str, object]:
    ws = wb.create_sheet("Calc")
    write_headers(ws, calc_headers(), width=13)

    for row in range(layout.SIM_FIRST_ROW, LAST_SIM_ROW + 1):
        for offset, formula in enumerate(calc_inputs_for(row)):
            ws.cell(row=row, column=offset + 1, value=formula)
        write_both_branches(ws, row, first_col=1)

    final = len(INPUT_HEADERS) + 2 * len(CALC_STEPS) + 1
    weapon_stats = build_weapon_stat_calc(wb, start_col=final + 2)
    return {
        "att_rules": get_column_letter(INPUT_HEADERS.index("att_rules") + 1),
        "def_rules": get_column_letter(INPUT_HEADERS.index("def_rules") + 1),
        "hits": get_column_letter(final),
        "wounds": get_column_letter(final + 1),
        "weapon_stats": weapon_stats,
    }


# --------------------------------------------------------------------------
# simulation tab
# --------------------------------------------------------------------------

def combat_formula(row: int, w: Dict[str, str]) -> str:
    """Ranged/Melee, derived from whether the chosen weapon has a Range."""
    return f'=IF({sim("Weapon", row)}="","",IF({w["range"]}>0,"Ranged","Melee"))'


def weapon_stats_formula(row: int, w: Dict[str, str]) -> str:
    """A readable summary of the chosen weapon's own profile, e.g. 18", A1, AP(1), Shred."""
    rng, atk, ap = w["range"], w["attacks"], w["ap"]
    blast, deadly, impact, rules = w["blast"], w["deadly"], w["impact"], w["rules"]

    parts = [
        f'IF({rng}>0,{rng}&"""","")',
        f'"A"&{atk}',
        f'IF({ap}>0,"AP("&{ap}&")","")',
        f'IF({blast}>0,"Blast("&{blast}&")","")',
        f'IF({deadly}>0,"Deadly("&{deadly}&")","")',
        f'IF({impact}>0,"Impact("&{impact}&")","")',
        rules,
    ]
    return f'=IF({sim("Weapon", row)}="","",{textjoin(", ", True, *parts)})'


def build_simulation(wb: Workbook, ranges, state_ranges, calc_cols) -> None:
    ws = wb["Simulation"]
    write_headers(ws, layout.SIMULATION_COLUMNS, width=15)
    ws.column_dimensions[get_column_letter(layout.col_index("Attacker"))].width = 20
    ws.column_dimensions[get_column_letter(layout.col_index("Defender"))].width = 20
    ws.column_dimensions[get_column_letter(layout.col_index("Weapon stats"))].width = 26
    ws.column_dimensions[get_column_letter(layout.col_index("Att. Special Rules"))].width = 28
    ws.column_dimensions[get_column_letter(layout.col_index("Def. Special Rules"))].width = 28

    for offset, seeded in enumerate(csvdata.simulation()):
        row = layout.SIM_FIRST_ROW + offset
        if row > LAST_SIM_ROW:
            break
        for column, value in enumerate(seeded, start=1):
            if value != "":
                ws.cell(row=row, column=column, value=value)

    attacker = get_column_letter(layout.col_index("Attacker"))
    numeric_outputs = (
        (layout.col_index("Hits"), calc_cols["hits"]),
        (layout.col_index("Wounds"), calc_cols["wounds"]),
    )
    text_outputs = (
        (layout.col_index("Att. Special Rules"), calc_cols["att_rules"]),
        (layout.col_index("Def. Special Rules"), calc_cols["def_rules"]),
    )

    for row in range(layout.SIM_FIRST_ROW, LAST_SIM_ROW + 1):
        guard = f'IF(${attacker}{row}="","",'
        for column, letter in numeric_outputs:
            cell = ws.cell(row=row, column=column, value=f"={guard}Calc!${letter}{row})")
            cell.number_format = "0.00"
            cell.fill = OUTPUT_FILL
        for column, letter in text_outputs:
            cell = ws.cell(row=row, column=column, value=f"={guard}Calc!${letter}{row})")
            cell.alignment = Alignment(wrap_text=True)
            cell.fill = OUTPUT_FILL

        w = calc_cols["weapon_stats"][row]
        combat_cell = ws.cell(row=row, column=layout.col_index("Combat"), value=combat_formula(row, w))
        combat_cell.fill = OUTPUT_FILL
        stats_cell = ws.cell(
            row=row, column=layout.col_index("Weapon stats"), value=weapon_stats_formula(row, w)
        )
        stats_cell.alignment = Alignment(wrap_text=True)
        stats_cell.fill = OUTPUT_FILL

    # faction dropdowns
    for header in ("Att. faction", "Def. faction"):
        dv = list_validation(
            f"Factions!$A${SRC_FIRST}:$A${layout.last_row('Factions')}"
        )
        ws.add_data_validation(dv)
        letter = get_column_letter(layout.col_index(header))
        dv.add(f"{letter}{layout.SIM_FIRST_ROW}:{letter}{LAST_SIM_ROW}")

    # per-row cascading dropdowns
    column_for_label = {
        "Attacker units": "Attacker",
        "Attacker upgrades": "Att. upgrades",
        "Weapons": "Weapon",
        "Defender units": "Defender",
        "Defender upgrades": "Def. upgrades",
    }
    for row, per_label in ranges.items():
        for label, source in per_label.items():
            dv = list_validation(source)
            ws.add_data_validation(dv)
            letter = get_column_letter(layout.col_index(column_for_label[label]))
            dv.add(f"{letter}{row}")

    # per-row Melee state / Range >9" dropdown, filtered by the row's derived Combat
    state_letter = get_column_letter(layout.col_index(STATE_COL))
    for row, source in state_ranges.items():
        dv = list_validation(source)
        ws.add_data_validation(dv)
        dv.add(f"{state_letter}{row}")

    # fixed-option dropdowns
    for header, options in (
        ("In cover", layout.YES_NO),
        ("Variable effect", layout.VARIABLE_EFFECT_OPTIONS),
    ):
        dv = inline_validation(options)
        ws.add_data_validation(dv)
        letter = get_column_letter(layout.col_index(header))
        dv.add(f"{letter}{layout.SIM_FIRST_ROW}:{letter}{LAST_SIM_ROW}")


# --------------------------------------------------------------------------
# tests tab
# --------------------------------------------------------------------------

def scenario_inputs(s: Scenario) -> List:
    active, weight = variable_effect(s)
    return [
        s.models, s.attacks, s.quality, s.defense, s.def_models,
        s.ap, s.blast, s.deadly, s.impact, s.ravage, s.tough,
        ", ".join(sorted(s.att_rules)), ", ".join(sorted(s.def_rules)),
        1 if active else 0, float(weight),
        s.is_melee, s.is_charging,
        s.is_melee and s.melee_state == "Fatigued",
        s.long_range, s.in_cover,
    ]


def build_tests(wb: Workbook) -> None:
    """Carry the pytest scenarios into the sheet so the formulas self-check."""
    ws = wb.create_sheet("Tests")
    headers = (
        ["Scenario"]
        + calc_headers()
        + ["Expected hits", "Actual hits", "Expected wounds", "Actual wounds", "Result"]
    )
    write_headers(ws, headers, width=13)
    ws.column_dimensions["A"].width = 34

    for i, (name, scenario) in enumerate(TEST_SCENARIOS):
        row = 2 + i
        ws.cell(row=row, column=1, value=name)
        for offset, value in enumerate(scenario_inputs(scenario)):
            ws.cell(row=row, column=2 + offset, value=value)

        cells = write_both_branches(ws, row, first_col=2)
        expected = resolve(scenario)

        base = 1 + len(calc_headers()) + 1
        ws.cell(row=row, column=base, value=float(expected.hits))
        ws.cell(row=row, column=base + 1, value=f"={cells['hits']}")
        ws.cell(row=row, column=base + 2, value=float(expected.wounds))
        ws.cell(row=row, column=base + 3, value=f"={cells['wounds']}")

        eh, ah = get_column_letter(base), get_column_letter(base + 1)
        ew, aw = get_column_letter(base + 2), get_column_letter(base + 3)
        ws.cell(
            row=row,
            column=base + 4,
            value=(
                f"=IF(AND(ABS({ah}{row}-{eh}{row})<0.000001,"
                f"ABS({aw}{row}-{ew}{row})<0.000001),\"PASS\",\"FAIL\")"
            ),
        )

    summary = 2 + len(TEST_SCENARIOS)
    result_col = get_column_letter(len(headers))
    ws.cell(row=summary + 1, column=1, value="All scenarios")
    ws.cell(
        row=summary + 1,
        column=len(headers),
        value=(
            f'=IF(COUNTIF({result_col}2:{result_col}{summary - 1},"FAIL")=0,'
            f'"ALL PASS","FAILURES")'
        ),
    )


# --------------------------------------------------------------------------

def build_workbook() -> Workbook:
    wb = Workbook()
    wb.active.title = "Simulation"

    for title, columns, loader, options in DATA_TABS:
        build_data_tab(wb, title, columns, loader(), **options)
    ranges = build_lists(wb)
    state_ranges = build_state_lists(wb, start_column=layout.SIM_ROWS * len(LIST_SPECS) + 1)
    calc_cols = build_calc(wb)
    build_tests(wb)
    build_simulation(wb, ranges, state_ranges, calc_cols)

    for name in ("Lists", "Calc", "Tests"):
        wb[name].sheet_state = "hidden"
    return wb


def main() -> Path:
    out = Path("dist") / "opr_crunch.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    build_workbook().save(out)
    return out


if __name__ == "__main__":
    print(f"wrote {main()}")
